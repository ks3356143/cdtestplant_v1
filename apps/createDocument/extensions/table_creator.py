from pathlib import Path
from typing import Literal
from docxtpl import DocxTemplate
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.shared import Mm
from apps.createDocument.extensions.tools import set_table_border_by_cell_position, set_cell_margins
from utils.path_utils import project_path
from utils.chen_response import ChenResponse
from django.shortcuts import get_object_or_404
from apps.project.models import Project, Round
from django.db.models import QuerySet

# 创建当前轮次别名
RoundType = Literal["0", "not0", "last"]
chinese_round_name: list = ['一', '二', '三', '四', '五', '六', '七', '八', '九', '十']

# 通用生成静态软件项、静态硬件项、动态软件项、动态硬件信息、测评数据的context，包含fontnote和table
def create_table_context(table_data: list[list[str]], doc: DocxTemplate, rounds_map: list[list[str]] = None,
                         current_round: str = None):
    """
        注意：该函数会增加一列序号列，并且支持单元格内回车换行（段落换行）
        传入当前轮次以及rounds_map来过滤一些其他轮次的数据，为None则不过滤
    """
    # 过滤数据处理
    if rounds_map is not None and current_round is not None and current_round != 'last':
        filtered_data = []
        for i, row in enumerate(table_data):
            # 第一行作为表头，始终保留
            if i == 0:
                filtered_data.append(row)
            else:
                # 检查该行是否属于当前轮次
                if current_round in rounds_map[i]:
                    filtered_data.append(row)
        # table_data先替换后再执行下方生成表格
        table_data = filtered_data

    subdoc = doc.new_subdoc()
    rows = len(table_data)
    cols = len(table_data[0]) + 1
    table = subdoc.add_table(rows=rows, cols=cols)

    # 单元格处理
    for row in range(rows):
        for col in range(cols):
            cell = table.cell(row, col)  # 从上倒下，从左到右取cell
            set_cell_margins(cell, left=100, right=100, top=100, bottom=100)

            # 获取要显示的文本内容（字符串或按行拆分后的列表）
            if col == 0:
                # 序号列
                lines = ["序号"] if row == 0 else [str(row)]
            else:
                raw_text = table_data[row][col - 1]
                # 按换行符 \n 拆分为多个段落
                lines = raw_text.split('\n') if raw_text else ['']

            # 清空单元格原有段落（add_table 默认有一个段落）
            cell.text = ""
            # 删除默认段落，稍后统一添加
            for para in cell.paragraphs:
                p = para._element
                p.getparent().remove(p)

            # 逐个添加段落
            for i, line in enumerate(lines):
                if i == 0:
                    para = cell.add_paragraph(line)
                else:
                    para = cell.add_paragraph(line)

                # 设置段落对齐（第一列居中，其他左对齐，可根据需要调整）
                if col == 0:
                    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                else:
                    para.alignment = WD_ALIGN_PARAGRAPH.LEFT

                # 对第一行（表头）设置黑体字体
                if row == 0:
                    for run in para.runs:
                        run.font.name = '黑体'
                        run._element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
                        run.font.bold = False
                    # 表头段落居中（覆盖前面的 left）
                    para.alignment = WD_ALIGN_PARAGRAPH.CENTER

            # 垂直居中
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    # 设置序号列宽度
    for cell in table.columns[0].cells:
        cell.width = Mm(15)
        for para in cell.paragraphs:
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # 表格居中
    table.alignment = WD_ALIGN_PARAGRAPH.CENTER
    # 设置表格外边框
    set_table_border_by_cell_position(table)
    return subdoc

# 统一静态软件项、静态硬件项、动态软件项、动态硬件信息、测评数据5个的word生成 - 模版模式
def uniform_static_dynamic_response(id: int, filename: str, r_filename: str, model,
                                    current_round: str = "0", isHsm: bool = None) -> ChenResponse | None:
    """ 通过形参isHsm判断是否是回归测试说明 """
    project_obj = get_object_or_404(Project, id=id)
    input_path = Path.cwd() / 'media' / project_path(id) / 'form_template' / 'dg' / filename  # 取相同模版
    doc = DocxTemplate(input_path)
    qs = model.objects.filter(project=project_obj)
    if qs.exists():
        obj = qs.first()
        table_data = obj.table
        # 回归测试说明生成多轮次回归测试说明的5个接口分支
        if isHsm:
            hround_list: QuerySet[Round] = project_obj.pField.exclude(key='0')
            for hround in hround_list:
                round_key = hround.key
                cname = chinese_round_name[int(round_key)]  # 取中文：一、二、三...
                # key就是current_round，这里就解决文件名和那个的问题
                subdoc = create_table_context(table_data, doc, obj.rounds_map, round_key)
                context = {
                    'fontnote': obj.fontnote,
                    'table': subdoc,
                }
                doc.render(context, autoescape=True)
                try:
                    doc.save(Path.cwd() / "media" / project_path(id) / "output_dir/hsm" / "".join([f"第{cname}轮", r_filename]))
                except PermissionError as e:
                    return ChenResponse(status=400, code=400, message="模版文件已打开，请关闭后再试，{0}".format(e))
            return ChenResponse(status=200,code=200,message="多个轮次5接口渲染完毕，文档生成完毕")

        # 新增：传入rounds_map进行渲染
        subdoc = create_table_context(table_data, doc, obj.rounds_map, current_round)
        context = {
            'fontnote': obj.fontnote,
            'table': subdoc,
        }
        doc.render(context, autoescape=True)
        try:
            doc.save(Path.cwd() / "media" / project_path(id) / "output_dir" / r_filename)
            return ChenResponse(status=200, code=200, message="文档生成成功！")
        except PermissionError as e:
            return ChenResponse(status=400, code=400, message="模版文件已打开，请关闭后再试，{0}".format(e))
    return None
